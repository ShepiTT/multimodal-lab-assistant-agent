"""
文本解析服务 - 直接读取文本类文件
支持: txt, md, html, csv, xlsx
"""
import csv
from pathlib import Path
from typing import Optional


class TextParser:
    """直接读取文本类文件"""
    
    # 支持常见代码文件和配置文件
    SUPPORTED_EXTENSIONS = {'.txt', '.md', '.html', '.htm', '.csv', '.xlsx', '.py', '.js', '.java', '.c', '.cpp', '.h', '.json', '.xml', '.yaml', '.yml'}
    
    @classmethod
    def can_parse(cls, file_path: str) -> bool:
        ext = Path(file_path).suffix.lower()
        return ext in cls.SUPPORTED_EXTENSIONS
    
    @classmethod
    def parse(cls, file_path: str, encoding: str = 'utf-8') -> str:
        """解析文本文件，返回文本内容"""
        path = Path(file_path)
        ext = path.suffix.lower()
        
        if ext in {'.txt', '.md', '.html', '.htm'}:
            return cls._read_text(path, encoding)
        elif ext == '.csv':
            return cls._read_csv(path, encoding)
        elif ext == '.xlsx':
            return cls._read_xlsx(path)
        elif ext in {'.py', '.js', '.java', '.c', '.cpp', '.h', '.json', '.xml', '.yaml', '.yml'}:
            # 代码和配置文件，直接读取
            return cls._read_text(path, encoding)
        else:
            raise ValueError(f"不支持的文件类型: {ext}")
    
    @staticmethod
    def _read_text(path: Path, encoding: str) -> str:
        """读取纯文本文件"""
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            # 尝试其他编码
            for enc in ['gbk', 'gb2312', 'latin-1']:
                try:
                    return path.read_text(encoding=enc)
                except UnicodeDecodeError:
                    continue
            raise ValueError("无法识别文件编码")
    
    @staticmethod
    def _read_csv(path: Path, encoding: str) -> str:
        """读取 CSV 文件，转为表格文本"""
        rows = []
        try:
            with path.open('r', encoding=encoding, newline='') as f:
                reader = csv.reader(f)
                for row in reader:
                    rows.append(' | '.join(row))
        except UnicodeDecodeError:
            with path.open('r', encoding='gbk', newline='') as f:
                reader = csv.reader(f)
                for row in reader:
                    rows.append(' | '.join(row))
        return '\n'.join(rows)
    
    @staticmethod
    def _read_xlsx(path: Path) -> str:
        """读取 Excel 文件"""
        try:
            import openpyxl
        except ImportError:
            raise ImportError("请安装 openpyxl: pip install openpyxl")
        
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        result = []
        
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            result.append(f"=== Sheet: {sheet_name} ===")
            for row in sheet.iter_rows(values_only=True):
                row_text = ' | '.join(str(cell) if cell is not None else '' for cell in row)
                if row_text.strip():
                    result.append(row_text)
        
        wb.close()
        return '\n'.join(result)
